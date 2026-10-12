#!/usr/bin/env python3
"""So sheet MO mới với catalog.json theo NỘI DUNG (cho chế độ update-catalog).

  sheet_diff.py "<file.xlsx>"                       # chỉ báo cáo: mới / sửa chữ / đổi số dòng / bị xoá
  sheet_diff.py "<file.xlsx>" --apply --snapshot 2026-11-01 [--images]

Cần `openpyxl` (và `Pillow` nếu dùng --images): `python3 -m pip install --user openpyxl pillow`.

--apply:
  - mục đổi số dòng → đổi ID (ID cũ vào `old_ids`), đổi luôn heading `## MO-x` trong references/,
    `depends_on`, `duplicate_of`, `recipe`;
  - mục sửa chữ → cập nhật dữ liệu sheet, GIỮ phần recipe; đánh dấu `needs_review`;
  - mục mới → thêm khung (group "misc", status active, `needs_review`) — người bảo trì phải viết recipe;
  - mục không còn trong sheet → `status: deprecated` (không xoá).
Luôn chạy lại `catalog.py check` + `catalog.py render` sau khi apply.
"""
import argparse
import datetime
import difflib
import io
import json
import os
import re
import sys

SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CATALOG = os.path.join(SKILL_DIR, "catalog.json")
URL_RE = re.compile(r"https?://[^\s)]+")
EST_DEV = re.compile(r"Estimate\s*Dev\s*:\s*([0-9.,]+)\s*(h|phút)?", re.I)
EST_QC = re.compile(r"Estimate\s*QC\s*:\s*([0-9.,]+)\s*h?\s*(\(([^)]*)\))?", re.I)

try:
    import openpyxl
except ImportError:
    sys.exit("Thiếu openpyxl: python3 -m pip install --user openpyxl")


def norm(text):
    import unicodedata
    text = URL_RE.sub(" ", str(text or "")).lower().replace("đ", "d")
    text = re.sub(r"estimate\s*(dev|qc)\s*:[^\n]*", " ", text)
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def clean(text):
    text = URL_RE.sub("", str(text or ""))
    text = re.sub(r"Estimate\s*(Dev|QC)\s*:[^\n]*", "", text, flags=re.I)
    return re.sub(r"\s+", " ", text).strip(" :-\n")


def merged_lookup(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        tl = ws.cell(rng.min_row, rng.min_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = (tl, rng)
    return m


def cell(ws, merged, r, c):
    hit = merged.get((r, c))
    v = (hit[0] if hit else ws.cell(r, c)).value
    if isinstance(v, str):
        v = v.replace("\r", "").strip() or None
    return v


def links(ws, merged, r, cols):
    figma, docs = [], []
    for c in cols:
        hit = merged.get((r, c))
        cl = hit[0] if hit else ws.cell(r, c)
        targets = [cl.hyperlink.target] if cl.hyperlink is not None and cl.hyperlink.target else []
        targets += URL_RE.findall(str(cl.value or ""))
        for t in targets:
            (figma if "figma.com" in t else docs).append(t)
    return list(dict.fromkeys(figma)), list(dict.fromkeys(docs))


def estimates(*texts):
    blob = "\n".join(str(t) for t in texts if t)
    out = {"dev_h": None, "qc_h": None, "qc_note": None}
    m = EST_DEV.search(blob)
    if m:
        n = float(m.group(1).replace(",", "."))
        out["dev_h"] = round(n / 60, 2) if (m.group(2) or "").lower() == "phút" else n
    m = EST_QC.search(blob)
    if m:
        out["qc_h"] = float(m.group(1).replace(",", "."))
        out["qc_note"] = m.group(3)
    return out


def read_sheet(path, sheet, app_cols):
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f"Không có sheet '{sheet}'. Có: {', '.join(wb.sheetnames)}")
    ws = wb[sheet]
    merged = merged_lookup(ws)
    header = {str(ws.cell(1, c).value).strip().lower(): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    rows = []
    for r in range(2, ws.max_row + 1):
        comp, detail = cell(ws, merged, r, 2), cell(ws, merged, r, 3)
        if not comp and not detail:
            continue
        multi = (r, 2) in merged and merged[(r, 2)][1].max_row > merged[(r, 2)][1].min_row
        c_comp, c_detail = clean(comp), clean(detail)
        title = f"{c_comp} — {c_detail}" if multi and c_detail and c_comp else (c_comp or c_detail)
        figma, docs = links(ws, merged, r, (2, 3, 8))
        status = {}
        for app in app_cols:
            c = header.get(app)
            if c and ws.cell(r, c).value is not None:
                status[app] = ws.cell(r, c).value
        rows.append({
            "row": r, "title": title[:160], "flow": clean(cell(ws, merged, r, 1)) or None,
            "sheet_text": {"component": comp, "detail": detail}, "owner": cell(ws, merged, r, 4),
            "needs_dev": ws.cell(r, 5).value, "needs_qc": ws.cell(r, 6).value,
            "est": estimates(comp, detail), "figma": figma, "docs": docs,
            "sheet_hint": cell(ws, merged, r, 8), "sheet_notes": cell(ws, merged, r, 25),
            "sheet_status": status,
            "sig": norm(f"{comp or ''} {detail or ''}"),
        })
    return wb, ws, rows


def match(rows, items):
    by_sig = {}
    for it in items:
        st = it.get("sheet_text") or {}
        by_sig.setdefault(norm(f"{st.get('component') or ''} {st.get('detail') or ''}"), []).append(it)
    pairs, used_items, used_rows = [], set(), set()
    for row in rows:
        cands = [it for it in by_sig.get(row["sig"], []) if it["id"] not in used_items]
        if cands:
            best = min(cands, key=lambda it: abs(it["row"] - row["row"]))
            pairs.append((row, best, 1.0))
            used_items.add(best["id"])
            used_rows.add(row["row"])
    rest_items = [it for it in items if it["id"] not in used_items]
    for row in rows:
        if row["row"] in used_rows:
            continue
        best, best_r = None, 0.0
        for it in rest_items:
            if it["id"] in used_items:
                continue
            st = it.get("sheet_text") or {}
            ratio = difflib.SequenceMatcher(None, row["sig"], norm(f"{st.get('component') or ''} {st.get('detail') or ''}")).ratio()
            if ratio > best_r:
                best, best_r = it, ratio
        if best is not None and best_r >= 0.75:
            pairs.append((row, best, best_r))
            used_items.add(best["id"])
            used_rows.add(row["row"])
    new_rows = [r for r in rows if r["row"] not in used_rows]
    removed = [it for it in items if it["id"] not in used_items]
    return pairs, new_rows, removed


def rename_in_references(mapping):
    if not mapping:
        return
    ref_dir = os.path.join(SKILL_DIR, "references")
    for fn in os.listdir(ref_dir):
        if not fn.endswith(".md"):
            continue
        p = os.path.join(ref_dir, fn)
        text = open(p).read()
        orig = text
        for old, new in mapping.items():
            text = re.sub(r"\b" + re.escape(old) + r"\b(?!\d)", f"@@{new}@@", text)
        text = re.sub(r"@@(MO-\d+)@@", r"\1", text)
        if text != orig:
            open(p, "w").write(text)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--snapshot", default=datetime.date.today().isoformat())
    ap.add_argument("--images", action="store_true", help="trích lại ảnh cột 'Hình ảnh' vào assets/refs")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()

    cat = json.load(open(CATALOG))
    meta = cat["meta"]
    sheet = meta["master_sheet"]
    wb, ws, rows = read_sheet(a.xlsx, sheet, meta.get("sheet_status_apps", []))
    items = [it for it in cat["items"] if it.get("sheet") == sheet and it.get("status") != "deprecated"]
    pairs, new_rows, removed = match(rows, items)

    moved = [(row, it) for row, it, r in pairs if row["row"] != it["row"]]
    changed = [(row, it, r) for row, it, r in pairs if r < 1.0]
    report = {
        "sheet": sheet, "rows_in_sheet": len(rows), "catalog_items": len(items),
        "moved": [{"old": it["id"], "new": f"MO-{row['row']}", "title": row["title"]} for row, it in moved],
        "changed_text": [{"id": it["id"], "new_id": f"MO-{row['row']}", "similarity": round(r, 2),
                          "old": (it.get("sheet_text") or {}), "new": row["sheet_text"]} for row, it, r in changed],
        "new": [{"id": f"MO-{row['row']}", "title": row["title"]} for row in new_rows],
        "removed": [{"id": it["id"], "title": it["title"]} for it in removed],
    }
    if a.json:
        print(json.dumps(report, ensure_ascii=False, indent=1, default=str))
    else:
        print(f"Sheet '{sheet}': {len(rows)} dòng · catalog: {len(items)} mục")
        for key, label in (("moved", "ĐỔI SỐ DÒNG"), ("changed_text", "SỬA CHỮ"), ("new", "MỤC MỚI"),
                           ("removed", "KHÔNG CÒN TRONG SHEET")):
            print(f"\n## {label}: {len(report[key])}")
            for x in report[key]:
                if key == "moved":
                    print(f"  {x['old']} → {x['new']}  {x['title'][:80]}")
                elif key == "changed_text":
                    print(f"  {x['id']} (→ {x['new_id']}, giống {x['similarity']}):")
                    print(f"     cũ: {norm(' '.join(str(v) for v in x['old'].values() if v))[:120]}")
                    print(f"     mới: {norm(' '.join(str(v) for v in x['new'].values() if v))[:120]}")
                else:
                    print(f"  {x['id']}  {x['title'][:90]}")
    if not a.apply:
        print("\n(Chỉ báo cáo. Trình user duyệt rồi chạy lại với --apply.)")
        return

    mapping = {it["id"]: f"MO-{row['row']}" for row, it in moved}
    for row, it, ratio in pairs:
        new_id = f"MO-{row['row']}"
        if it["id"] != new_id:
            it.setdefault("old_ids", []).append(it["id"])
        it["id"], it["row"] = new_id, row["row"]
        for k in ("flow", "sheet_text", "owner", "needs_dev", "needs_qc", "est", "figma", "docs", "sheet_hint",
                  "sheet_notes", "sheet_status"):
            it[k] = row[k]
        if ratio < 1.0:
            it["needs_review"] = f"Sheet sửa chữ ({a.snapshot}) — soát lại requirement/recipe"
    for it in removed:
        it["status"] = "deprecated"
        it["deprecated_reason"] = f"Không còn trong sheet '{sheet}' bản {a.snapshot}"
    for row in new_rows:
        base = {k: row[k] for k in ("row", "title", "flow", "sheet_text", "owner", "needs_dev", "needs_qc", "est",
                                    "figma", "docs", "sheet_hint", "sheet_notes", "sheet_status")}
        base.update({"id": f"MO-{row['row']}", "sheet": sheet, "group": "misc", "status": "active", "images": [],
                     "old_ids": [], "needs_review": f"Mục mới ({a.snapshot}) — chọn group, viết recipe",
                     "aliases": [], "requirement": "", "acceptance": [], "sdk": None,
                     "canonical": {"app": meta.get("default_ref", "a013"), "files": [], "symbols": [], "verified_commit": meta.get("verified_commit")},
                     "rc_keys": [], "events": [], "placements": [], "depends_on": [],
                     "detect": {"any": [], "all": []}, "recipe": "", "qc": []})
        cat["items"].append(base)
    for it in cat["items"]:
        it["depends_on"] = [mapping.get(d, d) for d in it.get("depends_on") or []]
        if it.get("duplicate_of") in mapping:
            it["duplicate_of"] = mapping[it["duplicate_of"]]
        if it.get("recipe"):
            it["recipe"] = re.sub(r"§\s*(MO-\d+)", lambda m: "§ " + mapping.get(m.group(1), m.group(1)), it["recipe"])
    rename_in_references(mapping)
    if a.images:
        refs = os.path.join(SKILL_DIR, "assets", "refs")
        for fn in os.listdir(refs):
            if fn.startswith("MO-"):
                os.remove(os.path.join(refs, fn))
        from PIL import Image
        by_row = {}
        for im in sorted(getattr(ws, "_images", []), key=lambda i: (i.anchor._from.row, i.anchor._from.col)):
            r = im.anchor._from.row + 1
            n = len(by_row.setdefault(r, [])) + 1
            name = f"MO-{r}-{n}.jpg"
            pic = Image.open(io.BytesIO(im._data())).convert("RGB")
            pic.thumbnail((1600, 1600))
            pic.save(os.path.join(refs, name), "JPEG", quality=80, optimize=True)
            by_row[r].append(f"assets/refs/{name}")
        for it in cat["items"]:
            if it.get("sheet") == sheet:
                it["images"] = by_row.get(it["row"], [])
    cat["items"].sort(key=lambda it: ({"Checklist - MO New": 0}.get(it.get("sheet"), 1), it.get("sheet", ""), it["row"]))
    meta["snapshot"] = a.snapshot
    with open(CATALOG, "w") as f:
        json.dump(cat, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print(f"\nĐã cập nhật catalog.json (snapshot {a.snapshot}). Chạy: catalog.py check && catalog.py render")


if __name__ == "__main__":
    main()
